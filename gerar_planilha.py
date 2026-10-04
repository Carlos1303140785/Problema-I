"""Gera Planilha.xlsx e disciplinas.js a partir da matriz do PPC TADS 2026 (cap. 3.2 e Fig. 3).

Uso: python gerar_planilha.py          (cria os arquivos; recusa se Planilha.xlsx já existir)
     python gerar_planilha.py --forcar  (recria tudo; APAGA projetos e professores já preenchidos)
"""

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

if os.path.exists("Planilha.xlsx") and "--forcar" not in sys.argv:
    sys.exit("Planilha.xlsx já existe. Rodar de novo apagaria os projetos enviados. Use --forcar se tiver certeza.")


# (id, sigla, nome, semestre, créditos, [pré-requisitos])
D = [
    ("MatDisc", "SI220", "Matemática Discreta", 1, 4, []),
    ("Prog1", "SI100", "Algoritmos e Programação de Computadores I", 1, 4, []),
    ("LabFerrProg", "SI103", "Lab. de Ferramentas de Programação", 1, 4, []),
    ("Arquit", "TT106", "Organização e Arquitetura de Computadores", 1, 4, []),
    ("AtividAlgorit", "SI104", "Atividades Práticas em Algoritmos", 1, 2, []),
    ("Estatis", "ST211", "Estatística", 2, 4, []),
    ("Analise1", "SI305", "Análise de Sistemas de Informação I", 2, 4, ["Prog1"]),
    ("EngSoft1", "SI206", "Engenharia de Software I", 2, 4, ["Prog1"]),
    ("CienTecSoci", "SI205", "Ciência, Tecnologia e Sociedade", 2, 2, []),
    ("Prog2", "SI204", "Algoritmos e Programação de Computadores II", 2, 2, ["Prog1"]),
    ("MatCienDad", "SI306", "Matemática para Ciência de Dados", 3, 4, []),
    ("POO1", "SI300", "Programação Orientada a Objetos I", 3, 4, ["Prog1"]),
    ("EngSoft2", "SI304", "Engenharia de Software II", 3, 4, ["EngSoft1"]),
    ("BD1", "ST567", "Banco de Dados I", 3, 4, ["Prog1"]),
    ("IntroIHC", "SI404", "Introdução a Interfaces Humano-Computador", 3, 2, ["EngSoft1"]),
    ("ED1", "SI201", "Estrutura de Dados I", 4, 4, ["Prog1"]),
    ("POO2", "SI400", "Programação Orientada a Objetos II", 4, 4, ["POO1"]),
    ("SO", "TT304", "Sistemas Operacionais", 4, 4, ["Arquit"]),
    ("BD2", "ST767", "Banco de Dados II", 4, 4, ["BD1"]),
    ("AtivIHC", "SI406", "Atividades Práticas em Interação Humano-Computador", 4, 2, ["IntroIHC"]),
    ("Web", "SI401", "Programação para a Web", 4, 4, ["POO1"]),
    ("Redes", "ST568", "Redes de Comunicação I", 5, 4, ["Arquit"]),
    ("Analise2", "SI405", "Análise de Sistemas de Informação II", 5, 4, ["Analise1"]),
    ("GestProj", "TT060", "Gestão de Projetos", 5, 4, ["EngSoft1"]),
    ("IA", "SI702", "Inteligência Artificial", 5, 4, ["ED1", "MatCienDad", "Prog2"]),
    ("ED2", "SI010", "Estrutura de Dados II", 5, 4, ["ED1"]),
    ("ProjInt", "SI503", "Projeto Integrador", 5, 4, ["POO1", "BD2", "EngSoft2", "IntroIHC", "Analise1", "ED1"]),
    ("DispMov", "SI700", "Programação para Dispositivos Móveis", 6, 4, ["POO2", "ED1"]),
    ("AprendMaqui", "SI602", "Introdução ao Aprendizado de Máquina", 6, 4, ["MatCienDad", "Prog2", "Estatis"]),
    ("Empreend", "SI800", "Empreendedorismo e Inovação", 6, 2, []),
    ("ProParaDi", "SI603", "Processamento Paralelo e Distribuído", 6, 4, ["ED1", "Redes", "SO", "Arquit"]),
    ("AtvCoEx", "SI919", "Atividades Complementares de Extensão", 6, 4, []),
    ("AtvCo", "SI920", "Atividades Complementares", 6, 12, []),
]

H = 15  # 1 crédito = 15 h (2010 h = 134 créditos)

wb = Workbook()
wb.remove(wb.worksheets[0])  # remove a aba "Sheet" que vem por padrão
ws = wb.create_sheet("Disciplinas")

cab = [
    "ID",
    "Sigla",
    "Nome",
    "Semestre",
    "Créditos",
    "Carga horária (h)",
    "Professores",
    "Pré-requisitos (IDs separados por vírgula)",
]

ws.append(cab)

for i, s, n, sem, c, p in D:
    r = ws.max_row + 1
    ws.append([i, s, n, sem, c, f"=E{r}*{H}", "", ", ".join(p)])

pj = wb.create_sheet("Projetos")
pj.append(["Data", "Nome", "Curso", "Semestre", "Disciplina (ID)", "GitHub"])

lg = wb.create_sheet("LEIA-ME")

for linha in [
    "Disciplinas: preencha só a coluna G (Professores, separados por vírgula). Pré-requisitos vieram da Figura 3 do PPC.",
    "Carga horária = créditos x 15 h (PPC: 2010 h = 134 créditos).",
    "Projetos: preenchida automaticamente pelo formulário (parser.py). Ex.: 2026-10-02 | Ana Souza | TADS | 2026-1 | POO1 | https://github.com/ana/projeto",
    "Disciplinas livres (QualqDisc) não aparecem no currículo interativo.",
]:
    lg.append([linha])

for planilha in (ws, pj):
    for celula in planilha[1]:
        celula.font = Font(
            name="Arial",
            bold=True,
            color="FFFFFF",
        )
        celula.fill = PatternFill(
            "solid",
            fgColor="44475A",
        )
        celula.alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )

for planilha, larguras in (
    (ws, [14, 8, 48, 9, 9, 12, 36, 40]),
    (pj, [12, 26, 12, 12, 16, 50]),
):
    for indice, largura in enumerate(larguras):
        planilha.column_dimensions[chr(65 + indice)].width = largura

    for row in planilha.iter_rows(min_row=2):
        for celula in row:
            celula.font = Font(name="Arial")

lg.column_dimensions["A"].width = 130

wb.save("Planilha.xlsx")

js = [
    {
        "id": i,
        "sigla": s,
        "nome": n,
        "semestre": sem,
        "creditos": c,
        "carga": c * H,
        "professores": [],
        "prereqs": p,
    }
    for i, s, n, sem, c, p in D
]

with open("disciplinas.js", "w", encoding="utf-8") as arquivo:
    arquivo.write(
        "const DISCIPLINAS = "
        + json.dumps(js, ensure_ascii=False, indent=1)
        + ";\n"
    )

print(len(D), "disciplinas")
