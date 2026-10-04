"""Servidor local simples: liga o site à Planilha.xlsx.

Uso:  python parser.py        (requer: pip install openpyxl)

GET  /api/disciplinas -> lê a aba Disciplinas
GET  /api/projetos    -> lê a aba Projetos
POST /api/projetos    -> acrescenta uma linha na aba Projetos
"""

import datetime
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import load_workbook

PLANILHA = "Planilha.xlsx"


def lista(valor):
    return [x.strip() for x in str(valor or "").split(",") if x.strip()]


def calcular_carga(horas, creditos):
    if horas:
        return horas
    return creditos * 15 if isinstance(creditos, (int, float)) else None


def ler_disciplinas():
    ws = load_workbook(PLANILHA, data_only=True)["Disciplinas"]
    return [
        {
            "id": i, "sigla": s, "nome": n, "semestre": sem, "creditos": c,
            "carga": calcular_carga(h, c), "professores": lista(prof), "prereqs": lista(pre),
        }
        for i, s, n, sem, c, h, prof, pre in ws.iter_rows(min_row=2, values_only=True)
        if i
    ]


def ler_projetos():
    ws = load_workbook(PLANILHA, data_only=True)["Projetos"]
    return [
        {"nome": n, "curso": c, "semestre": s, "disciplina": d, "github": g}
        for _, n, c, s, d, g in ws.iter_rows(min_row=2, values_only=True)
        if g
    ]


def salvar_projeto(p):
    wb = load_workbook(PLANILHA)
    wb["Projetos"].append([
        datetime.datetime.now(tz=datetime.timezone.utc).astimezone().date().isoformat(),
        p.get("nome"), p.get("curso"), p.get("semestre"),
        p.get("disciplina"), p.get("github"),
    ])
    wb.save(PLANILHA)


class Handler(BaseHTTPRequestHandler):
    def responder(self, codigo, corpo=None):
        self.send_response(codigo)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        if corpo is None:
            self.end_headers()
            return
        dados = json.dumps(corpo, ensure_ascii=False).encode()
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(dados)))
        self.end_headers()
        self.wfile.write(dados)

    def do_OPTIONS(self):
        self.responder(204)

    def do_GET(self):
        if self.path == "/api/disciplinas":
            self.responder(200, ler_disciplinas())
        elif self.path == "/api/projetos":
            self.responder(200, ler_projetos())
        else:
            self.responder(404, {"erro": "não encontrado"})

    def do_POST(self):
        if self.path != "/api/projetos":
            return self.responder(404, {"erro": "não encontrado"})
        try:
            tamanho = int(self.headers.get("Content-Length", 0))
            salvar_projeto(json.loads(self.rfile.read(tamanho)))
            self.responder(201, {"ok": True})
        except PermissionError:
            self.responder(500, {"erro": "Feche a planilha no Excel e tente de novo"})
        except OSError:
            self.responder(500, {"erro": "Não foi possível acessar a planilha"})
        except (ValueError, AttributeError) as e:
            self.responder(400, {"erro": str(e)})


if __name__ == "__main__":
    print("Parser em http://localhost:8000")
    ThreadingHTTPServer(("", 8000), Handler).serve_forever()
